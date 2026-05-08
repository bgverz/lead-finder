"""Tests for sales workflow features: contact quality, reveal priority, status tracking,
feedback import, profile comparison, and new Excel sheets."""

from __future__ import annotations

import csv
from pathlib import Path

import pytest

from lead_pipeline.models import Lead
from lead_pipeline.output.excel import CRM_EXPORT_COLUMNS, OUTPUT_COLUMNS, ExcelExporter
from lead_pipeline.scoring.scorer import LeadScorer
from lead_pipeline.utils.config import load_config
from lead_pipeline.utils.database import LeadDatabase, VALID_STATUSES


# ---------------------------------------------------------------------------
# contact_quality
# ---------------------------------------------------------------------------

def test_contact_quality_full_when_email_and_phone():
    lead = Lead(
        name="Full Contact",
        property_value=0,
        property_address="",
        email="test@example.com",
        phone="+12035550101",
        phone_valid=True,
        apollo_obfuscated=False,
    )
    assert lead.contact_quality == "full"


def test_contact_quality_partial_when_only_email():
    lead = Lead(
        name="Email Only",
        property_value=0,
        property_address="",
        email="test@example.com",
        phone="",
        phone_valid=False,
        apollo_obfuscated=False,
    )
    assert lead.contact_quality == "partial"


def test_contact_quality_partial_when_only_phone():
    lead = Lead(
        name="Phone Only",
        property_value=0,
        property_address="",
        email="",
        phone="+12035550101",
        phone_valid=True,
        apollo_obfuscated=False,
    )
    assert lead.contact_quality == "partial"


def test_contact_quality_partial_from_apollo_flags():
    lead = Lead(
        name="Apollo Flags",
        property_value=0,
        property_address="",
        apollo_person_id="p1",
        apollo_has_email=True,
        apollo_has_direct_phone=True,
        apollo_obfuscated=True,
    )
    assert lead.contact_quality == "partial"


def test_contact_quality_none_when_no_info():
    lead = Lead(name="No Contact", property_value=0, property_address="")
    assert lead.contact_quality == "none"


def test_contact_quality_obfuscated_no_flags_is_none():
    lead = Lead(
        name="Obfuscated",
        property_value=0,
        property_address="",
        email="obfuscated@example.com",
        phone="+10000000000",
        phone_valid=True,
        apollo_obfuscated=True,
        # No apollo_has_email / apollo_has_direct_phone flags set.
    )
    # Obfuscated contacts with no availability flags → no usable contact info.
    assert lead.contact_quality == "none"


# ---------------------------------------------------------------------------
# contact_reveal_priority
# ---------------------------------------------------------------------------

def _scored_apollo_lead(**overrides) -> Lead:
    config = load_config("config/config.example.yaml")
    scorer = LeadScorer(config)
    defaults = dict(
        name="Test",
        property_value=0,
        property_address="",
        apollo_person_id="p1",
        apollo_title="CEO",
        apollo_obfuscated=True,
    )
    lead = Lead(**{**defaults, **overrides})
    return scorer.score(lead)


def test_reveal_priority_skip_when_not_obfuscated():
    lead = Lead(
        name="Clear",
        property_value=0,
        property_address="",
        apollo_person_id="p1",
        apollo_obfuscated=False,
        lead_score=80,
    )
    assert lead.contact_reveal_priority == "skip"


def test_reveal_priority_skip_when_no_apollo_id():
    lead = Lead(name="No Apollo", property_value=3_000_000, property_address="1 Main", lead_score=85)
    assert lead.contact_reveal_priority == "skip"


def test_reveal_priority_skip_when_score_below_40():
    lead = Lead(
        name="Low Score",
        property_value=0,
        property_address="",
        apollo_person_id="p1",
        apollo_obfuscated=True,
        lead_score=35,
    )
    assert lead.contact_reveal_priority == "skip"


def test_reveal_priority_high_for_top_title_good_score():
    lead = _scored_apollo_lead(apollo_title="Founder", apollo_organization="Acme Capital")
    # Founder(35) + org(5) + email flags depend on setup; just check it's high enough
    lead.lead_score = 70
    assert lead.contact_reveal_priority == "high"


def test_reveal_priority_medium_for_moderate_score():
    lead = Lead(
        name="Med",
        property_value=0,
        property_address="",
        apollo_person_id="p1",
        apollo_title="Partner",
        apollo_obfuscated=True,
        lead_score=52,
    )
    assert lead.contact_reveal_priority == "medium"


def test_reveal_priority_low_for_apollo_with_low_score():
    lead = Lead(
        name="Low Rev",
        property_value=0,
        property_address="",
        apollo_person_id="p1",
        apollo_title="Analyst",
        apollo_obfuscated=True,
        lead_score=41,
    )
    assert lead.contact_reveal_priority == "low"


def test_reveal_priority_skip_when_contact_already_full():
    lead = Lead(
        name="Has Contact",
        property_value=0,
        property_address="",
        apollo_person_id="p1",
        apollo_obfuscated=False,
        email="test@example.com",
        phone="+12035550101",
        phone_valid=True,
        lead_score=80,
    )
    assert lead.contact_reveal_priority == "skip"


# ---------------------------------------------------------------------------
# lead_status field
# ---------------------------------------------------------------------------

def test_lead_status_default_is_new():
    lead = Lead(name="New Lead", property_value=0, property_address="")
    assert lead.lead_status == "new"


def test_lead_status_in_to_output_dict():
    lead = Lead(name="A", property_value=0, property_address="", lead_status="contacted")
    assert lead.to_output_dict()["lead_status"] == "contacted"


def test_contact_quality_in_to_output_dict():
    lead = Lead(
        name="A",
        property_value=0,
        property_address="",
        email="a@b.com",
        phone="+1",
        phone_valid=True,
    )
    d = lead.to_output_dict()
    assert "contact_quality" in d
    assert "contact_reveal_priority" in d
    assert d["contact_quality"] == "full"


def test_apollo_search_profile_in_to_output_dict():
    lead = Lead(
        name="A", property_value=0, property_address="", apollo_search_profile="vc_pe"
    )
    assert lead.to_output_dict()["apollo_search_profile"] == "vc_pe"


# ---------------------------------------------------------------------------
# Database — status tracking
# ---------------------------------------------------------------------------

def test_upsert_and_apply_stored_status(tmp_path):
    db = LeadDatabase(str(tmp_path / "test.db"))
    lead = Lead(name="Jordan", property_value=3_000_000, property_address="1 Main St")
    db.upsert_many([lead])

    ok = db.update_lead_status("Jordan", "1 Main St", "contacted", "Left voicemail")
    assert ok is True

    fresh_lead = Lead(name="Jordan", property_value=3_000_000, property_address="1 Main St")
    db.apply_stored_statuses([fresh_lead])
    assert fresh_lead.lead_status == "contacted"
    assert fresh_lead.status_notes == "Left voicemail"
    db.close()


def test_upsert_preserves_non_new_status(tmp_path):
    db = LeadDatabase(str(tmp_path / "test.db"))
    lead = Lead(name="Sam", property_value=2_000_000, property_address="2 Oak Ave")
    db.upsert_many([lead])
    db.update_lead_status("Sam", "2 Oak Ave", "converted")

    # Re-upsert with status="new" should not overwrite "converted".
    lead2 = Lead(name="Sam", property_value=2_500_000, property_address="2 Oak Ave")
    db.upsert_many([lead2])

    row = db.conn.execute(
        "SELECT lead_status FROM leads WHERE name='Sam'"
    ).fetchone()
    assert row["lead_status"] == "converted"
    db.close()


def test_update_lead_status_returns_false_if_not_found(tmp_path):
    db = LeadDatabase(str(tmp_path / "test.db"))
    result = db.update_lead_status("Nobody", "Nowhere", "contacted")
    assert result is False
    db.close()


def test_load_reveal_candidates_excludes_do_not_contact(tmp_path):
    db = LeadDatabase(str(tmp_path / "test.db"))
    keep = Lead(name="Keep", property_value=0, property_address="1 A", lead_score=70, score_tier="Hot")
    skip = Lead(name="Skip", property_value=0, property_address="2 B", lead_score=80, score_tier="Hot")
    db.upsert_many([keep, skip])
    db.update_lead_status("Skip", "2 B", "do_not_contact")

    candidates = db.load_reveal_candidates(min_score=60.0)
    names = [r["name"] for r in candidates]
    assert "Keep" in names
    assert "Skip" not in names
    db.close()


def test_load_profile_stats_aggregates_correctly(tmp_path):
    db = LeadDatabase(str(tmp_path / "test.db"))
    leads = [
        Lead(name="A", property_value=0, property_address="1",
             lead_score=85, score_tier="Hot", apollo_search_profile="vc_pe"),
        Lead(name="B", property_value=0, property_address="2",
             lead_score=65, score_tier="Warm", apollo_search_profile="vc_pe"),
        Lead(name="C", property_value=0, property_address="3",
             lead_score=45, score_tier="Cool", apollo_search_profile="founders"),
    ]
    db.upsert_many(leads)
    stats = db.load_profile_stats()

    vc_pe = next((s for s in stats if s["profile"] == "vc_pe"), None)
    assert vc_pe is not None
    assert vc_pe["total"] == 2
    assert vc_pe["hot"] == 1
    assert vc_pe["warm"] == 1

    founders = next((s for s in stats if s["profile"] == "founders"), None)
    assert founders is not None
    assert founders["total"] == 1
    db.close()


# ---------------------------------------------------------------------------
# Database — import_feedback
# ---------------------------------------------------------------------------

def test_import_feedback_updates_matching_rows(tmp_path):
    db = LeadDatabase(str(tmp_path / "test.db"))
    lead = Lead(name="Alice", property_value=0, property_address="1 A")
    db.upsert_many([lead])

    updated, not_found = db.import_feedback([
        {"name": "Alice", "status": "contacted", "notes": "Email sent", "property_address": "1 A"},
    ])
    assert updated == 1
    assert not_found == 0

    row = db.conn.execute("SELECT lead_status, status_notes FROM leads WHERE name='Alice'").fetchone()
    assert row["lead_status"] == "contacted"
    assert row["status_notes"] == "Email sent"
    db.close()


def test_import_feedback_counts_not_found(tmp_path):
    db = LeadDatabase(str(tmp_path / "test.db"))
    _, not_found = db.import_feedback([
        {"name": "Ghost", "status": "skipped"},
    ])
    assert not_found == 1
    db.close()


def test_import_feedback_matches_by_name_only_when_no_address(tmp_path):
    db = LeadDatabase(str(tmp_path / "test.db"))
    lead = Lead(name="Bob", property_value=0, property_address="5 Elm St")
    db.upsert_many([lead])

    updated, _ = db.import_feedback([{"name": "Bob", "status": "revealed"}])
    assert updated == 1
    db.close()


# ---------------------------------------------------------------------------
# CLI — import-feedback CSV parsing
# ---------------------------------------------------------------------------

def test_import_feedback_cli_dry_run(tmp_path):
    """Dry-run should not write to the database."""
    from click.testing import CliRunner
    from lead_pipeline.cli import import_feedback

    db_path = tmp_path / "test.db"
    db = LeadDatabase(str(db_path))
    lead = Lead(name="Carol", property_value=0, property_address="7 Pine Ave")
    db.upsert_many([lead])
    db.close()

    csv_path = tmp_path / "feedback.csv"
    with open(csv_path, "w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=["name", "status", "notes"])
        writer.writeheader()
        writer.writerow({"name": "Carol", "status": "contacted", "notes": "Called"})

    runner = CliRunner()
    result = runner.invoke(
        import_feedback,
        ["--config", "config/config.example.yaml", "--dry-run", str(csv_path)],
    )
    assert result.exit_code == 0
    assert "dry-run" in result.output.lower()

    # Status must not have changed.
    db2 = LeadDatabase(str(db_path))
    row = db2.conn.execute("SELECT lead_status FROM leads WHERE name='Carol'").fetchone()
    assert row["lead_status"] == "new"
    db2.close()


def test_import_feedback_cli_rejects_invalid_status(tmp_path):
    from click.testing import CliRunner
    from lead_pipeline.cli import import_feedback

    csv_path = tmp_path / "bad.csv"
    with open(csv_path, "w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=["name", "status"])
        writer.writeheader()
        writer.writerow({"name": "Dave", "status": "called_back"})  # invalid

    runner = CliRunner()
    result = runner.invoke(
        import_feedback,
        ["--config", "config/config.example.yaml", "--dry-run", str(csv_path)],
    )
    assert result.exit_code != 0 or "invalid" in result.output.lower()


# ---------------------------------------------------------------------------
# Valid statuses constant
# ---------------------------------------------------------------------------

def test_valid_statuses_set():
    assert "new" in VALID_STATUSES
    assert "contacted" in VALID_STATUSES
    assert "revealed" in VALID_STATUSES
    assert "converted" in VALID_STATUSES
    assert "do_not_contact" in VALID_STATUSES
    assert "skipped" in VALID_STATUSES


# ---------------------------------------------------------------------------
# Excel — new columns and sheets
# ---------------------------------------------------------------------------

def test_output_columns_has_new_workflow_fields():
    for col in ("contact_quality", "contact_reveal_priority", "lead_status", "status_notes", "apollo_search_profile"):
        assert col in OUTPUT_COLUMNS, f"Missing column: {col}"


def test_output_columns_preserves_existing_positions():
    # Guard against accidentally shifting columns the existing test checks.
    assert OUTPUT_COLUMNS[0] == "name"
    assert OUTPUT_COLUMNS[1] == "lead_score"
    assert OUTPUT_COLUMNS[2] == "score_tier"
    assert OUTPUT_COLUMNS[3] == "next_action"
    assert OUTPUT_COLUMNS[4] == "lead_summary"
    assert OUTPUT_COLUMNS[15:18] == ["property_lookup_status", "property_lookup_reason", "address_needed"]


def test_crm_export_columns_contain_key_fields():
    for col in ("name", "lead_score", "contact_quality", "contact_reveal_priority",
                "lead_status", "phone", "email", "apollo_title", "apollo_organization"):
        assert col in CRM_EXPORT_COLUMNS, f"Missing CRM column: {col}"


def test_excel_export_creates_new_sheets(tmp_path):
    from openpyxl import load_workbook

    config = load_config("config/config.example.yaml")
    config.output.output_dir = str(tmp_path)
    exporter = ExcelExporter(config)

    leads = [
        Lead(
            name="High Rev",
            property_value=0,
            property_address="",
            apollo_person_id="p1",
            apollo_title="CEO",
            apollo_obfuscated=True,
            lead_score=70,
            score_tier="Hot",
        ),
        Lead(
            name="CRM Ready",
            property_value=2_000_000,
            property_address="1 Main",
            email="crm@example.com",
            phone="+12035550101",
            phone_valid=True,
            lead_score=65,
            score_tier="Warm",
        ),
    ]
    # Set reveal priority manually so the sheet filter works.
    leads[0].lead_score = 70

    path = exporter.export(leads, tmp_path / "test.xlsx")
    wb = load_workbook(path)

    assert "Reveal Candidates" in wb.sheetnames
    assert "CRM Ready" in wb.sheetnames
    assert "Export Ready" in wb.sheetnames


def test_reveal_candidates_sheet_has_correct_headers(tmp_path):
    from openpyxl import load_workbook

    config = load_config("config/config.example.yaml")
    config.output.output_dir = str(tmp_path)
    exporter = ExcelExporter(config)

    path = exporter.export([], tmp_path / "empty.xlsx")
    wb = load_workbook(path)

    # Reveal Candidates uses OUTPUT_COLUMNS.
    headers = [cell.value for cell in wb["Reveal Candidates"][1]]
    assert headers == OUTPUT_COLUMNS


def test_export_ready_sheet_uses_crm_columns(tmp_path):
    from openpyxl import load_workbook

    config = load_config("config/config.example.yaml")
    config.output.output_dir = str(tmp_path)
    exporter = ExcelExporter(config)

    path = exporter.export([], tmp_path / "empty2.xlsx")
    wb = load_workbook(path)

    headers = [cell.value for cell in wb["Export Ready"][1]]
    assert headers == CRM_EXPORT_COLUMNS
