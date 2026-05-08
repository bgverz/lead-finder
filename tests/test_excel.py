from pathlib import Path

from openpyxl import load_workbook

from lead_pipeline.models import Lead
from lead_pipeline.output.excel import OUTPUT_COLUMNS, ExcelExporter
from lead_pipeline.utils.config import load_config


def test_excel_export_creates_tier_sheets(tmp_path: Path):
    config = load_config("config/config.example.yaml")
    config.output.output_dir = str(tmp_path)
    exporter = ExcelExporter(config)
    lead = Lead(
        name="Priya Shah",
        property_value=2_850_000,
        property_address="88 Lake Avenue",
        area_income=162_000,
        lead_score=60,
        score_tier="Warm",
    )

    output = exporter.export([lead], tmp_path / "leads.xlsx")
    workbook = load_workbook(output)

    assert {"All Leads", "Hot", "Warm", "Cool", "Skip"}.issubset(workbook.sheetnames)
    headers = [cell.value for cell in workbook["All Leads"][1]]
    assert headers == OUTPUT_COLUMNS
    # Spot-check key columns are present and in correct order.
    assert headers[0] == "name"
    assert headers[1] == "lead_score"
    assert headers[2] == "score_tier"
    assert headers[3] == "next_action"
    assert headers[4] == "lead_summary"
    assert headers[15:18] == ["property_lookup_status", "property_lookup_reason", "address_needed"]
    assert "apollo_signal_score" in headers
    assert workbook["All Leads"].freeze_panes == "A2"
    # property_value column should be currency-formatted (now at index 9).
    prop_col_letter = chr(ord("A") + OUTPUT_COLUMNS.index("property_value"))
    assert workbook["All Leads"][f"{prop_col_letter}2"].number_format == "$#,##0"
